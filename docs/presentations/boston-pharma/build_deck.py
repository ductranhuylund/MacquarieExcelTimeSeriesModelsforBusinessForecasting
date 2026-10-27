from pathlib import Path
import textwrap
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor

OUT=Path(__file__).resolve().parent
OUT.mkdir(parents=True, exist_ok=True)
W,H=960,540
T='#087F78'; D='#073E43'; M='#DFF2ED'; BG='#F5F9F7'; INK='#163B40'; GRAY='#587371'; WHITE='#FFFFFF'; GOLD='#C89340'
pdfmetrics.registerFont(TTFont('DV','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'))
pdfmetrics.registerFont(TTFont('DVB','/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'))
slides=[]; cur=[]; notes=[]
def box(x,y,w,h,fill,radius=False):cur.append(('box',x,y,w,h,fill,radius))
def photo(path,x,y,w,h):cur.append(('image',x,y,w,h,str(path)))
def txt(x,y,w,text,size=18,color=INK,bold=False,leading=None):
    leading=leading or size*1.38
    lines=[]
    for paragraph in text.split('\n'):
        if not paragraph:lines.append('');continue
        line=''
        for word in paragraph.split():
            trial=(line+' '+word).strip()
            if pdfmetrics.stringWidth(trial,'DVB' if bold else 'DV',size)>w and line:
                lines.append(line);line=word
            else:line=trial
        lines.append(line)
    for j,line in enumerate(lines):cur.append(('text',x,y+j*leading,w,line,size,color,bold))
    return len(lines)*leading

def start(section,title,sub='',dark=False):
    global cur
    cur=[];slides.append(cur)
    box(0,0,W,H,D if dark else BG)
    box(0,0,9,H,T)
    txt(40,22,600,'BOSTON PHARMA',12,M if dark else T,True)
    txt(40,54,880,section.upper(),10,M if dark else GRAY,True)
    txt(40,80,880,title,30,WHITE if dark else D,True,36)
    if sub:txt(40,130,880,sub,12,M if dark else GRAY)
    txt(40,510,810,'PHÊ DUYỆT PO  /  SO SÁNH PHƯƠNG ÁN  /  27.10.2026',8,M if dark else GRAY)
    txt(900,510,40,f'{len(slides):02d}',9,M if dark else T,True)

def card(x,y,w,h,kicker,title,body,accent=T):
    box(x,y,w,h,WHITE,True);box(x,y,4,h,accent)
    txt(x+18,y+17,w-36,kicker,10,accent,True)
    hh=txt(x+18,y+42,w-36,title,21,D,True,27)
    txt(x+18,y+55+hh,w-36,body,14,GRAY,False,21)

def band(text,y=458):
    box(40,y,880,35,M,True);txt(54,y+8,850,text,11,T,True)

def step(x,y,w,num,title,body):
    box(x,y,w,183,WHITE,True);box(x+15,y+16,30,27,T,True)
    txt(x+22,y+20,24,str(num),12,WHITE,True)
    txt(x+15,y+58,w-30,title,18,D,True,23)
    txt(x+15,y+113,w-30,body,12,GRAY,False,18)

start('Đề xuất giải pháp','Hai hướng tiếp cận.\nMột quyết định phê duyệt tin cậy.',dark=True)
txt(42,193,630,'Đối chiếu bằng Python\nhay phê duyệt qua Reply Email?',23,WHITE,False,32)
box(42,298,540,83,T,True)
txt(60,313,510,'PYTHON WORKSPACE  ×  SAP EMAIL APPROVAL',14,WHITE,True)
txt(60,344,510,'PO: F011 • F012 • F013  |  PR: giai đoạn sau',12,M)
for i,(a,b) in enumerate([('01','Đúng PO'),('02','Đúng thẩm quyền'),('03','Đúng dấu vết')]):
 txt(696,206+i*70,70,a,24,M,True);txt(742,211+i*70,175,b,15,WHITE)
notes.append('Thiết kế theo màu xanh ngọc của ảnh nhận diện Boston Pharma người dùng cung cấp; không có bộ brand guideline hoặc file logo gốc trong gói tài liệu.')

start('01 / Kết luận điều hành','Hai phương án giải quyết hai nhu cầu khác nhau','Cần chọn theo cách làm việc của người duyệt và phạm vi tích hợp mong muốn.')
card(40,174,428,239,'PYTHON • ĐÃ CÓ MÃ NGUỒN','Đối chiếu trên một màn hình','Dành cho người cần xem Base và SAP cạnh nhau, kiểm tra chi tiết PO rồi duyệt thủ công trên từng hệ thống.')
card(488,174,432,239,'EMAIL • THIẾT KẾ ĐỀ XUẤT','Duyệt từ hộp thư','Dành cho người cần phản hồi từ điện thoại hoặc máy tính; backend SAP xác thực và thực thi quyết định.',GOLD)
band('Khuyến nghị: hoàn thiện đối chiếu hiện tại; thử nghiệm email ở phạm vi hẹp trước khi mở rộng.')
notes.append('Nguồn: README và mã Python trong repository; tài liệu DOCX được cung cấp. Phương án email chưa có bằng chứng triển khai trong nguồn đầu vào.')

start('02 / Phạm vi thực tế','Phân biệt phần đang có và phần cần xây','Không coi bản mô tả luồng email là tính năng đã sẵn sàng vận hành.')
card(40,175,428,265,'HIỆN TẠI','Python Workspace','• Base API: danh sách F011 / F012 / F013, lọc trạng thái và cache trong phiên.\n• Website Base + SAP GUI nhúng.\n• Tìm PO bằng ME29N; duyệt do người dùng thực hiện.')
card(488,175,432,265,'CẦN TRIỂN KHAI','Reply Email','• Trigger gửi thư, nhận và phân tích thư trả lời.\n• Xác thực, kiểm tra quyền và release code.\n• Thực thi SAP, phản hồi kết quả và lưu nhật ký.',GOLD)
band('Python: SAP GUI còn cần UAT trên Windows. Email: chưa mô tả đồng bộ kết quả về Base.')
notes.append('App không tự Release/Save. PR bị khóa chờ mapping. Kiểm thử tự động hiện dùng mô phỏng COM; không thay thế UAT SAP Windows.')

start('03 / Phương án A','Python: người duyệt kiểm soát hai hệ thống','Đăng nhập SAP chỉ tạo phiên. Nút “Tìm nhiệm vụ SAP” mới mở ME29N và tìm PO.')
for i,(a,b) in enumerate([('Kết nối phiên','Token API Base, website Base và đăng nhập SAP.'),('Chọn job Base','Chọn workflow và nhiệm vụ; lấy PO cần đối chiếu.'),('Tìm PO SAP','Mở ME29N, điền PO và kiểm tra PO đang hiển thị.'),('Duyệt & tiếp theo','Người dùng duyệt từng bên rồi chuyển nhiệm vụ.')]):
 step(40+i*225,188,205,i+1,a,b)
 if i<3:txt(248+i*225,263,22,'→',20,T,True)
band('Không tự xác nhận cả hai bên đã duyệt. Đổi job Base không đồng nghĩa SAP đã đổi PO.')
notes.append('Luồng theo mã nguồn base_sap_workspace.py tại commit 705933d. PO từ tiêu đề là gợi ý văn bản; có trường hợp phải chọn hoặc nhập PO.')

start('04 / Trải nghiệm A','Một không gian đối chiếu cho người xử lý PO','Minh họa thiết kế, không phải ảnh hệ thống trực tiếp.')
box(40,176,880,248,WHITE,True)
box(52,188,352,34,M,True);txt(64,197,330,'BASE  /  F011     F012     F013',12,T,True)
box(550,188,358,34,M,True);txt(562,197,330,'SAP  /  ME29N',12,T,True)
txt(68,244,318,'Danh sách nhiệm vụ',15,D,True)
txt(68,279,318,'Job được chọn → PO yêu cầu\nNội dung • tài liệu • trao đổi\nDuyệt trực tiếp trên website Base',13,GRAY,False,25)
box(417,266,119,62,T,True);txt(429,278,97,'Tìm PO\nSAP →',13,WHITE,True,21)
txt(566,244,322,'PO đang hiển thị',15,D,True)
txt(566,279,322,'Nhà cung cấp • giá trị • điều khoản\nKiểm tra đúng PO trước khi duyệt\nDuyệt trực tiếp trong SAP GUI',13,GRAY,False,25)
band('Điểm kiểm soát quan trọng: tách “PO yêu cầu” khỏi “PO thực tế đang hiển thị”.')
notes.append('Sơ đồ minh họa bằng shape, không chứa dữ liệu nghiệp vụ hoặc ảnh màn hình thật.')

start('Ảnh chụp app / 01','Bắt đầu phiên: kết nối trước, xử lý sau','Ảnh chụp Qt app hiện tại; dữ liệu và dịch vụ được mô phỏng, không có kết nối SAP thật.')
photo(OUT/'screenshots/04_session_start_demo.png',58,168,386,316)
card(478,174,442,122,'BASE','Hai kết nối độc lập','API lấy danh sách; website dùng để xem và duyệt. Token API không thay thế đăng nhập website.')
card(478,310,442,133,'SAP','Đăng nhập chưa mở ME29N','Chỉ khi bấm “Tìm nhiệm vụ SAP”, app mới mở transaction và tìm PO của job đã chọn.')
notes.append('Ảnh được capture từ WorkSessionDialog thực trong app, dùng test harness mô phỏng Base/SAP. Không phải ảnh SAP Windows trực tiếp. Tài liệu và sơ đồ gốc vẫn giữ nguyên phạm vi so sánh.')

start('Ảnh chụp app / 02','Hàng đợi Base: chọn workflow và nhiệm vụ','Ảnh chụp widget thật của app; mã PO và nội dung là dữ liệu mẫu.')
box(40,174,880,308,WHITE,True)
photo(OUT/'screenshots/02_base_queue_demo.png',54,188,618,258)
for y,k,title,body in [(190,'01','Chọn workflow','F011 / F012 / F013'),(277,'02','Lọc và tìm kiếm','Theo trạng thái, mã job, PO hoặc tiêu đề.'),(376,'03','Chuyển nhiệm vụ','Chọn dòng hoặc dùng Trước / Tiếp theo.')]:
 txt(697,y,200,k+' / '+title,12,T,True)
 txt(697,y+30,195,body,11,GRAY,False,17)
notes.append('Ảnh chụp vùng hàng đợi của Window.left bằng QWidget.grab. Dữ liệu mẫu, API và browser mô phỏng. Bảng giữ đầy đủ nhiệm vụ; 3 dòng chỉ là vùng hiển thị mặc định.')

start('Ảnh chụp app / 03','Đối chiếu PO yêu cầu với PO đang hiển thị','Ảnh chụp phần thông tin đối chiếu; ví dụ trùng PO được mô phỏng để giải thích giao diện.')
box(40,173,880,229,WHITE,True)
photo(OUT/'screenshots/03_sap_match_demo.png',55,183,850,207)
box(40,418,429,70,M,True)
txt(55,429,400,'PO yêu cầu từ Base',14,T,True)
txt(55,456,400,'Đầu vào của lượt tìm kiếm SAP.',12,GRAY)
box(487,418,433,70,M,True)
txt(502,429,405,'PO thực tế trên SAP',14,T,True)
txt(502,456,405,'Phải khớp trước khi người dùng duyệt.',12,GRAY)
notes.append('Chỉ ảnh widget thông tin của app; giá trị PO thực tế được giả lập, không phải đọc từ SAP trực tiếp.')

start('05 / Phương án B','Email: quyết định đi từ hộp thư vào SAP','Tóm tắt theo tài liệu “MÔ TẢ LUỒNG VẬN HÀNH (FLOW) CƠ CHẾ REPLY EMAIL”.')
items=[('SAP phát thư','PO đủ điều kiện duyệt; thư có tóm tắt, PDF và Ref Token.'),('Người duyệt Reply','APPROVE / OK hoặc REJECT; lý do ở dòng tiếp theo.'),('Backend kiểm tra','Ghép Ref Token, người gửi, quyền duyệt và trạng thái PO.'),('SAP thực thi','Gọi cơ chế release phù hợp; gửi mail kết quả sau xác nhận.')]
for i,(a,b) in enumerate(items):
 step(40+i*225,188,205,i+1,a,b)
 if i<3:txt(248+i*225,263,22,'→',20,T,True)
band('Luồng này tập trung vào SAP PO. Trạng thái job Base không tự đổi nếu chưa xây tích hợp bổ sung.')
notes.append('DOCX nêu SCOT/SO28 và BAPI_PO_RELEASE. Cần đội SAP xác minh kênh inbound, release strategy và API phù hợp phiên bản; sơ đồ ở đây là kiến trúc logic.')

start('06 / So sánh','Chọn theo nhu cầu, không chỉ theo số nút bấm','Đánh giá định tính từ phạm vi hiện có và thiết kế email được cung cấp.')
rows=[('Nơi thao tác','Windows • Base + SAP cạnh nhau','Hộp thư trên máy tính / điện thoại'),('Xem chi tiết & đối chiếu','Mạnh: xem trực tiếp cả hai hệ thống','Tóm tắt + PDF; có thể cần mở SAP'),('Thực hiện duyệt','Thủ công trên Base và SAP','Backend SAP thực hiện từ phản hồi'),('Phụ thuộc kỹ thuật','Base API, Qt, SAP GUI Scripting','Mail inbound/outbound, ABAP, phân quyền'),('Ảnh hưởng trạng thái Base','Người dùng thao tác trực tiếp','Cần thiết kế đồng bộ riêng'),('Mức sẵn sàng','Đã có mã; SAP cần UAT thực tế','Thiết kế; cần phát triển và kiểm thử'),('Phù hợp nhất','Người xử lý cần kiểm tra sâu','Người duyệt cần linh hoạt địa điểm')]
xs=[40,252,585]; widths=[212,333,335]
for x,w,t in zip(xs,widths,['TIÊU CHÍ','PYTHON WORKSPACE','REPLY EMAIL']):box(x,172,w,32,T);txt(x+12,181,w-24,t,11,WHITE,True)
for i,row in enumerate(rows):
 y=204+i*34
 for x,w,t in zip(xs,widths,row):box(x,y,w-1,33,WHITE if i%2==0 else M);txt(x+12,y+9,w-24,t,10,INK,i==0)
notes.append('Không có dữ liệu đo thời gian, chi phí hoặc tỷ lệ lỗi thực tế; không trình bày ưu thế định lượng chưa được đo.')

start('07 / Điều kiện triển khai email','Bốn điều kiện cần xác minh trước triển khai','Các khẳng định trong tài liệu là mục tiêu thiết kế; cần chứng minh bằng cấu hình và kiểm thử.')
card(40,174,428,126,'01 / DANH TÍNH','Email gửi về ≠ đủ xác thực','Ghép SU01 cần đi cùng xác thực kênh thư, token và quyền release thực tế.')
card(488,174,432,126,'02 / DUYỆT TRÙNG','Cần khóa và xử lý nguyên tử','Kiểm tra trạng thái trước khi ghi chưa đủ nếu email và SAP GUI duyệt đồng thời.')
card(40,316,428,126,'03 / NHẬT KÝ','Không mặc định là user của sếp','User chạy backend có thể là tài khoản kỹ thuật; cần lưu người yêu cầu và người thực thi.')
card(488,316,432,126,'04 / REJECT & THÀNH CÔNG','Cần chốt nghiệp vụ riêng','BAPI release không thay thế thiết kế từ chối; chỉ báo thành công sau commit và đọc lại.')
notes.append('DOCX dùng các diễn đạt “an toàn tuyệt đối”, không dump và chỉ một vết user sếp. Deck không xác nhận các điều đó khi chưa có mã, cấu hình và bằng chứng kiểm thử. Không hướng dẫn ghi thẳng EKKO/EKPO.')

start('08 / Kiểm soát vận hành','Mỗi phương án có một bộ kiểm soát trọng tâm','Cùng giữ trạng thái gốc của Base và SAP; không tạo cảm giác đã duyệt khi mới gửi yêu cầu.')
card(40,175,428,267,'PYTHON','Đúng màn hình, đúng PO','1. Xác minh PO SAP khớp PO yêu cầu.\n2. Dừng ở hộp thoại chưa nhận diện.\n3. Làm mới danh sách để kiểm tra kết quả.\n4. Ghi nhận trường hợp Base/SAP lệch trạng thái; không tự coi đã hoàn tất.')
card(488,175,432,267,'EMAIL','Đúng người, đúng phiên bản','1. Token gắn PO, cấp duyệt và phiên bản.\n2. Hết hạn, dùng một lần và chống phát lại.\n3. Kiểm tra quyền, khóa xử lý đồng thời.\n4. Đọc phần trả lời mới; không lấy APPROVE từ thư trích dẫn.',GOLD)
notes.append('Các kiểm soát email là đề xuất bổ sung, không phải toàn bộ đều có sẵn trong tài liệu. App hiện tại có kiểm tra PO, giữ hộp thoại lạ; chưa tự xác nhận cả hai hệ thống hoàn tất.')

start('09 / Các tình huống phải kiểm thử','Kiểm thử cả ngoại lệ và luồng thành công','Một bộ tình huống chung giúp so sánh công bằng hai phương án.')
rows=[('PO đúng / đủ quyền','Mở đúng PO; người dùng duyệt','Đúng cấp duyệt; commit rồi phản hồi'),('PO đã duyệt trên SAP','Đọc trạng thái, tránh thao tác thừa','Bỏ qua yêu cầu trùng; phản hồi đã xử lý'),('PO đổi sau khi phát mail','Xem dữ liệu hiện tại trên SAP','Từ chối token cũ hoặc yêu cầu duyệt lại'),('Lệch quyền / sai người','SAP chặn theo quyền người dùng','Không release; lưu lý do và báo lỗi'),('Lỗi mạng / timeout','Không mặc định đã thành công','Đối soát trạng thái trước khi retry'),('Base và SAP lệch nhau','Cần người xử lý đối chiếu','Cần cơ chế đồng bộ hoặc hàng đợi xử lý')]
for x,w,t in [(40,246,'TÌNH HUỐNG'),(286,310,'PYTHON'),(596,324,'EMAIL')]:box(x,176,w,33,T);txt(x+12,185,w-24,t,11,WHITE,True)
for i,row in enumerate(rows):
 y=209+i*39
 for x,w,t in zip([40,286,596],[246,310,324],row):box(x,y,w-1,38,WHITE if i%2==0 else M);txt(x+12,y+8,w-24,t,10,INK,False,14)
notes.append('Đây là tiêu chí kiểm thử đề xuất, không phải kết quả thử nghiệm. Trạng thái “đã duyệt” phải đánh giá theo release code/cấp duyệt và phiên bản chứng từ liên quan.')

start('10 / Lộ trình','Hai nhánh thử nghiệm, một bộ tiêu chí nghiệm thu','Chưa chốt lịch hoặc chi phí khi chưa khảo sát cấu hình SAP và hạ tầng email.')
for i,(a,b) in enumerate([('Chốt nghiệp vụ','Phạm vi PO, cấp duyệt, REJECT và trách nhiệm cập nhật Base.'),('Ổn định Python','UAT SAP Windows; kiểm tra đúng PO, popup, phiên và lỗi kết nối.'),('PoC Email','Một nhóm duyệt và một luồng PO; ABAP, mail, token và audit.'),('Đo & quyết định','So sánh thời gian, lỗi, tải hỗ trợ và khả năng truy vết.')]):
 step(40+i*225,188,205,i+1,a,b)
 if i<3:txt(248+i*225,263,22,'→',20,T,True)
band('Đội phối hợp: Mua hàng + người duyệt + SAP functional/ABAP/Basis + quản trị mail + IT ứng dụng.')
notes.append('PoC/UAT và trình tự là đề xuất; không phải kế hoạch cam kết. PR cần phân tích riêng vì chưa có mapping trong app và DOCX tập trung PO.')

start('11 / Đo hiệu quả','Đo bằng dữ liệu của Boston Pharma','Thiết lập baseline trước khi pilot; chưa gán mục tiêu phần trăm khi chưa đo thực tế.')
metrics=[('Thời gian','Từ yêu cầu đến duyệt thành công; tách thời gian chờ và thao tác.'),('Độ chính xác','Số lần sai PO, sai cấp duyệt, trạng thái Base/SAP không khớp.'),('Độ ổn định','Tỷ lệ thao tác lỗi; retry, thư trùng và yêu cầu cần hỗ trợ.'),('Truy vết','Tỷ lệ quyết định truy được PO, phiên bản, người yêu cầu, kết quả.')]
for i,(a,b) in enumerate(metrics):card(40+(i%2)*448,174+(i//2)*140,432,125,f'KPI 0{i+1}',a,b)
notes.append('Không trình bày mức tiết kiệm, ngân sách hay SLA giả định. Cần thống nhất cách đo và ngưỡng đạt trước pilot.')

start('12 / Đề xuất quyết định','Giữ đối chiếu sâu. Thử nghiệm duyệt linh hoạt.',dark=True)
txt(42,166,590,'Python cho người cần xem đầy đủ Base–SAP.\nEmail cho người cần phê duyệt từ xa.',23,WHITE,True,34)
for i,(a,b) in enumerate([('LÀM NGAY','Hoàn thiện UAT app Python hiện có.'),('THỬ NGHIỆM','PoC email trên một phạm vi PO hẹp.'),('CHỐT TRƯỚC KHI MỞ RỘNG','Ai đồng bộ Base? Ai xử lý ngoại lệ và audit?')]):
 y=276+i*58;txt(42,y,260,a,10,M,True);txt(305,y,590,b,15,WHITE)
txt(42,465,860,'Hai kênh có thể cùng tồn tại khi dùng chung quy tắc quyền và kiểm soát duyệt trùng.',12,M)
notes.append('Khuyến nghị dựa trên nhu cầu đã trao đổi: app phục vụ đối chiếu; email bổ sung kênh quyết định SAP. Hai kênh chỉ nên cùng vận hành sau kiểm thử concurrency và audit.')

start('Nguồn & giả định','Cơ sở của bản so sánh','Thông tin được rà soát ngày 27.10.2026. Không bao gồm dữ liệu PO, token hoặc thông tin đăng nhập.')
card(40,177,880,113,'NGUỒN 01 / APP PYTHON','Repository đã phát triển','ductranhuylund/MacquarieExcelTimeSeriesModelsforBusinessForecasting\nbase-sap-workspace/ • mã nguồn và README • commit 705933d')
card(40,306,880,113,'NGUỒN 02 / THIẾT KẾ EMAIL','Tài liệu người dùng cung cấp','“MÔ TẢ LUỒNG VẬN HÀNH (FLOW) CƠ CHẾ REPLY EMAIL.docx”\nLuồng 5 bước; APPROVE / REJECT; Ref Token; xử lý yêu cầu trùng.')
txt(42,444,870,'Màu xanh ngọc theo ảnh nhận diện đã cung cấp; không thay thế brand guideline chính thức.\nCác đề xuất kiểm soát, UAT và lộ trình trong deck chưa phải tính năng email đã triển khai.',11,GRAY,False,18)
notes.append('GitHub: https://github.com/ductranhuylund/MacquarieExcelTimeSeriesModelsforBusinessForecasting/tree/main/base-sap-workspace . Tài liệu email là nguồn mô tả, không phải yêu cầu thực thi hoặc chứng cứ hệ thống đang hoạt động.')

prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
prs.core_properties.title='Boston Pharma | Python Workspace và Reply Email'
prs.core_properties.subject='So sánh hai phương án phê duyệt PO'
prs.core_properties.author='Boston Pharma'
prs.core_properties.keywords='PO, Base, SAP, ME29N, Python, Reply Email'
pdf=canvas.Canvas(str(OUT/'Boston_Pharma_So_sanh_Python_Email.pdf'),pagesize=(W,H))
pdf.setTitle('Boston Pharma — Python Workspace và Reply Email')
for idx,objects in enumerate(slides):
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    slide.notes_slide.notes_text_frame.text=notes[idx]
    for obj in objects:
        if obj[0]=='box':
            _,x,y,w,h,fill,rounded=obj
            shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h))
            if rounded:
                try:shape.adjustments[0]=0.08
                except:pass
            shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(fill[1:]);shape.line.fill.background()
            pdf.setFillColor(HexColor(fill));pdf.setStrokeColor(HexColor(fill))
            if rounded:pdf.roundRect(x,H-y-h,w,h,6,fill=1,stroke=0)
            else:pdf.rect(x,H-y-h,w,h,fill=1,stroke=0)
        elif obj[0]=='image':
            _,x,y,w,h,path=obj
            slide.shapes.add_picture(path,Pt(x),Pt(y),width=Pt(w),height=Pt(h))
            pdf.drawImage(path,x,H-y-h,width=w,height=h,mask='auto')
        else:
            _,x,y,w,line,size,color,bold=obj
            shape=slide.shapes.add_textbox(Pt(x),Pt(y),Pt(w+6),Pt(size*1.5))
            tf=shape.text_frame;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.word_wrap=False
            p=tf.paragraphs[0];p.space_after=Pt(0);p.space_before=Pt(0)
            r=p.add_run();r.text=line;r.font.name='Arial';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=RGBColor.from_string(color[1:])
            pdf.setFillColor(HexColor(color));pdf.setFont('DVB' if bold else 'DV',size);pdf.drawString(x,H-y-size*.94,line)
    pdf.showPage()
prs.save(OUT/'Boston_Pharma_So_sanh_Python_Email.pptx');pdf.save()
# Export review thumbnails from the PDF companion.
import fitz
from PIL import Image,ImageOps,ImageDraw
pdfdoc=fitz.open(OUT/'Boston_Pharma_So_sanh_Python_Email.pdf')
thumbs=[]
for i,page in enumerate(pdfdoc):
 pix=page.get_pixmap(matrix=fitz.Matrix(.5,.5));im=Image.frombytes('RGB',[pix.width,pix.height],pix.samples)
 im=ImageOps.expand(im,border=8,fill='white');thumbs.append(im)
contact=Image.new('RGB',(thumbs[0].width*3,thumbs[0].height*6),'#CBD9D4')
for i,im in enumerate(thumbs):contact.paste(im,((i%3)*im.width,(i//3)*im.height))
contact.save('/tmp/boston_deck_contact.png')
# Validate content bounds, page count and editable PowerPoint structure.
for n,objects in enumerate(slides,1):
 for o in objects:
  if o[0]=='text':
   assert o[2]+o[5]*1.2 <= H, (n,o)
assert len(Presentation(OUT/'Boston_Pharma_So_sanh_Python_Email.pptx').slides)==17
print('Generated',len(slides),'slides:',[(p.name,p.stat().st_size) for p in OUT.iterdir()])
